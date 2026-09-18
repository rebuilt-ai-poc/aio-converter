"""PDF-source and PDF merge converter tests."""
from __future__ import annotations

from pathlib import Path

import pytest

from app.converters.pdf import pdf_to_txt, pdf_to_markdown, pdf_to_docx, merge_pdfs


# ---------------------------------------------------------------------------
# PDF -> TXT
# ---------------------------------------------------------------------------
def test_pdf_to_txt_extracts_text(tmp_path: Path, fixtures_dir: Path) -> None:
    out = tmp_path / "out.txt"
    pdf_to_txt(fixtures_dir / "simple.pdf", out, {})
    text = out.read_text(encoding="utf-8")
    assert "Hello, World!" in text
    assert "simple test PDF" in text


def test_pdf_to_txt_preserves_unicode(tmp_path: Path, fixtures_dir: Path) -> None:
    out = tmp_path / "out.txt"
    pdf_to_txt(fixtures_dir / "simple.pdf", out, {})
    text = out.read_text(encoding="utf-8")
    # Latin-extended round-trips through PyMuPDF's default Helvetica.
    assert "cafe" in text or "café" in text
    assert "naive" in text or "naïve" in text


def test_pdf_to_txt_preserves_page_order(tmp_path: Path, fixtures_dir: Path) -> None:
    out = tmp_path / "out.txt"
    pdf_to_txt(fixtures_dir / "two-page.pdf", out, {})
    text = out.read_text(encoding="utf-8")
    i1 = text.find("TWO#1")
    i2 = text.find("TWO#2")
    assert 0 <= i1 < i2, text


def test_pdf_to_txt_page_break_option(tmp_path: Path, fixtures_dir: Path) -> None:
    out = tmp_path / "out.txt"
    pdf_to_txt(fixtures_dir / "two-page.pdf", out, {"preserve_page_breaks": True})
    assert "\f" in out.read_text(encoding="utf-8")


# ---------------------------------------------------------------------------
# PDF -> Markdown
# ---------------------------------------------------------------------------
def test_pdf_to_markdown_produces_content(tmp_path: Path, fixtures_dir: Path) -> None:
    out = tmp_path / "out.md"
    pdf_to_markdown(fixtures_dir / "headings.pdf", out, {})
    md = out.read_text(encoding="utf-8")
    assert md.strip(), "markdown should not be empty"
    assert "Chapter One" in md
    assert "Introduction" in md


# ---------------------------------------------------------------------------
# PDF -> DOCX
# ---------------------------------------------------------------------------
def test_pdf_to_docx_opens_and_has_content(tmp_path: Path, fixtures_dir: Path) -> None:
    from docx import Document

    out = tmp_path / "out.docx"
    pdf_to_docx(fixtures_dir / "headings.pdf", out, {})
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Chapter One" in text
    assert "Introduction" in text
    assert "paragraph in the body text" in text


def test_pdf_to_docx_infers_heading(tmp_path: Path, fixtures_dir: Path) -> None:
    from docx import Document

    out = tmp_path / "out.docx"
    pdf_to_docx(fixtures_dir / "headings.pdf", out, {})
    doc = Document(str(out))
    styles = [(p.text.strip(), p.style.name) for p in doc.paragraphs]
    # The 22pt "Chapter One" line should be flagged as a Heading of some level.
    assert any("Heading" in style for text, style in styles if text == "Chapter One"), styles


def test_pdf_to_docx_preserves_table_content(tmp_path: Path, fixtures_dir: Path) -> None:
    from docx import Document

    out = tmp_path / "out.docx"
    pdf_to_docx(fixtures_dir / "table.pdf", out, {})
    doc = Document(str(out))
    # Table content should either land in a real docx table or in paragraphs.
    hit = False
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                if "North" in cell.text:
                    hit = True
    if not hit:
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "North" in text and "South" in text, text
    # Column labels should appear somewhere either way.
    all_text = "\n".join(p.text for p in doc.paragraphs) + " ".join(
        cell.text for t in doc.tables for row in t.rows for cell in row.cells
    )
    assert "Q1" in all_text and "Q2" in all_text


# ---------------------------------------------------------------------------
# PDF merge
# ---------------------------------------------------------------------------
def test_merge_pdfs_page_count(tmp_path: Path, fixtures_dir: Path) -> None:
    import pymupdf

    out = tmp_path / "merged.pdf"
    merge_pdfs(
        [fixtures_dir / "two-page.pdf", fixtures_dir / "three-page.pdf"],
        out,
    )
    doc = pymupdf.open(out)
    try:
        assert doc.page_count == 5
    finally:
        doc.close()


def test_merge_pdfs_preserves_order(tmp_path: Path, fixtures_dir: Path) -> None:
    import pymupdf

    out = tmp_path / "merged.pdf"
    merge_pdfs(
        [fixtures_dir / "three-page.pdf", fixtures_dir / "two-page.pdf"],
        out,
    )
    doc = pymupdf.open(out)
    try:
        first_page = doc[0].get_text("text")
        last_page = doc[-1].get_text("text")
    finally:
        doc.close()
    assert "THREE#1" in first_page
    assert "TWO#2" in last_page
